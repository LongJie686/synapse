"""Document loaders for various file formats."""

from __future__ import annotations

import csv
import json
from typing import Any
from pathlib import Path

from synapse_core.rag import BaseDocumentLoader, Document


class TextLoader(BaseDocumentLoader):
    """Load plain text files."""

    async def load(self, source: str, **kwargs: Any) -> list[Document]:
        path = Path(source)
        content = path.read_text(encoding=kwargs.get("encoding", "utf-8"))
        return [Document(content=content, source=str(path), metadata={"format": "txt", "size": len(content)})]


class MarkdownLoader(BaseDocumentLoader):
    """Load Markdown files."""

    async def load(self, source: str, **kwargs: Any) -> list[Document]:
        path = Path(source)
        content = path.read_text(encoding=kwargs.get("encoding", "utf-8"))
        return [Document(content=content, source=str(path), metadata={"format": "markdown", "size": len(content)})]


class CSVLoader(BaseDocumentLoader):
    """Load CSV files, one document per row."""

    async def load(self, source: str, **kwargs: Any) -> list[Document]:
        path = Path(source)
        encoding = kwargs.get("encoding", "utf-8")
        delimiter = kwargs.get("delimiter", ",")
        documents = []
        with open(path, encoding=encoding, newline="") as f:
            reader = csv.DictReader(f, delimiter=delimiter)
            for i, row in enumerate(reader):
                content = json.dumps(row, ensure_ascii=False)
                documents.append(Document(content=content, source=str(path), metadata={"format": "csv", "row": i, **row}))
        return documents


class PDFLoader(BaseDocumentLoader):
    """Load PDF files using pypdf."""

    async def load(self, source: str, **kwargs: Any) -> list[Document]:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("pip install pypdf")

        reader = PdfReader(source)
        documents = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(Document(content=text, source=source, page=i, metadata={"format": "pdf", "page": i}))
        return documents


class DocxLoader(BaseDocumentLoader):
    """Load DOCX files using python-docx."""

    async def load(self, source: str, **kwargs: Any) -> list[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("pip install python-docx")

        doc = DocxDocument(source)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n\n".join(paragraphs)
        return [Document(content=content, source=source, metadata={"format": "docx"})]


class URLLoader(BaseDocumentLoader):
    """Load content from URLs."""

    async def load(self, source: str, **kwargs: Any) -> list[Document]:
        import re

        try:
            import httpx
        except ImportError:
            raise ImportError("pip install httpx")

        async with httpx.AsyncClient(timeout=30) as client:
            response = await client.get(source)
            response.raise_for_status()
            content = response.text

        if "<html" in content.lower():
            content = re.sub(r"<script[^>]*>[\s\S]*?</script>", "", content, flags=re.IGNORECASE)
            content = re.sub(r"<style[^>]*>[\s\S]*?</style>", "", content, flags=re.IGNORECASE)
            content = re.sub(r"<[^>]+>", " ", content)
            content = re.sub(r"\s+", " ", content).strip()

        return [Document(content=content, source=source, metadata={"format": "url"})]


class DirectoryLoader(BaseDocumentLoader):
    """Load all supported files from a directory recursively."""

    _EXTENSION_LOADERS: dict[str, type[BaseDocumentLoader]] = {
        ".txt": TextLoader,
        ".md": MarkdownLoader,
        ".csv": CSVLoader,
        ".pdf": PDFLoader,
        ".docx": DocxLoader,
    }

    async def load(self, source: str, **kwargs: Any) -> list[Document]:
        directory = Path(source)
        if not directory.is_dir():
            raise ValueError(f"Not a directory: {source}")

        documents = []
        for file_path in sorted(directory.rglob("*")):
            if file_path.is_file() and file_path.suffix in self._EXTENSION_LOADERS:
                loader = self._EXTENSION_LOADERS[file_path.suffix]()
                try:
                    docs = await loader.load(str(file_path), **kwargs)
                    documents.extend(docs)
                except Exception as e:
                    documents.append(Document(content=f"Failed to load: {e}", source=str(file_path), metadata={"error": str(e)}))
        return documents
